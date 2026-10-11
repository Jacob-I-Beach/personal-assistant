"""Read structure and render an actual page inside the same isolated environment."""
import csv
import json
import subprocess
from pathlib import Path
import assistant

request = json.loads(Path('/task/request.json').read_text())
source = Path(request['document_path']).resolve()
if not source.is_relative_to(Path('/inputs')) or not source.is_file():
    raise ValueError('Choose a file supplied to this execution.')
page_number = max(1, min(500, int(request.get('page', 1))))
suffix = source.suffix.lower()
summary = {'filename': source.name, 'source_bytes': source.stat().st_size, 'page_requested': page_number}
def cell_value(value):
    if value is None or isinstance(value, (bool,int,float)):
        return value
    return str(value)[:100]

if suffix in {'.xlsx', '.xlsm'}:
    import openpyxl
    book = openpyxl.load_workbook(source, read_only=True, data_only=False, keep_links=False)
    values = openpyxl.load_workbook(source, read_only=True, data_only=True, keep_links=False)
    sheets = []
    for sheet in list(book)[:4]:
        rows = []
        for row in sheet.iter_rows(max_row=min(sheet.max_row, 20), max_col=min(sheet.max_column, 8)):
            rows.append([{'cell': cell.coordinate, 'formula': str(cell.value)[:100] if cell.data_type == 'f' else None,
                          'value': cell_value(cell.value) if cell.data_type != 'f' else None,
                          'cached_value': cell_value(values[sheet.title][cell.coordinate].value) if cell.data_type == 'f' else None} for cell in row])
        sheets.append({'name': sheet.title, 'rows': sheet.max_row, 'columns': sheet.max_column, 'preview': rows})
    summary['sheets'] = sheets
    summary['coverage'] = 'Up to four sheets, 20 rows and eight columns per sheet; cell/formula previews are capped at 100 characters. Read the source with Python for full values. Missing cached values are not zero.'
    book.close(); values.close()
elif suffix == '.docx':
    from docx import Document
    doc = Document(source)
    summary['paragraphs'] = [p.text[:800] for p in doc.paragraphs[:60]]
    summary['tables'] = [[[cell.text[:200] for cell in row.cells[:8]] for row in table.rows[:15]] for table in doc.tables[:4]]
    summary['coverage'] = 'Paragraph and table previews are bounded. Inspect the rendered page for layout and embedded figures.'
elif suffix in {'.csv', '.tsv'}:
    with source.open(encoding='utf-8-sig', newline='') as stream:
        reader = csv.reader(stream, delimiter='\t' if suffix == '.tsv' else ',')
        summary['rows'] = [row for _, row in zip(range(40), reader)]
    summary['coverage'] = 'First 40 rows; original file remains available for computation.'

pdf = source if suffix == '.pdf' else None
if suffix in {'.docx', '.xlsx', '.xlsm', '.pptx', '.odt', '.ods', '.odp'}:
    profile = Path('/tmp/lo-profile/user'); profile.mkdir(parents=True, exist_ok=True)
    (profile/'registrymodifications.xcu').write_text('''<?xml version="1.0" encoding="UTF-8"?><oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item><item oor:path="/org.openoffice.Office.Common/Load"><prop oor:name="UpdateDocMode" oor:op="fuse"><value>0</value></prop></item></oor:items>''')
    converted = subprocess.run(['libreoffice', '-env:UserInstallation=file:///tmp/lo-profile', '--headless', '--convert-to', 'pdf', '--outdir', '/outputs', str(source)], capture_output=True, timeout=45)
    candidate = Path('/outputs')/(source.stem+'.pdf')
    if converted.returncode == 0 and candidate.exists():
        pdf = candidate
        summary['rendering_note'] = 'Office layout rendered with LibreOffice; pagination can differ from Microsoft Office.'
    else:
        summary['rendering_note'] = 'The Office renderer could not produce a page. Structured previews remain available.'

if pdf:
    import pypdfium2 as pdfium
    document = pdfium.PdfDocument(pdf)
    summary['page_count'] = len(document)
    if page_number > len(document): raise ValueError('The requested page is outside this document.')
    page = document[page_number - 1]
    text_page = page.get_textpage()
    summary['page_text'] = text_page.get_text_range()[:18000]
    summary['text_coverage'] = 'Text on the requested page, capped at 18,000 characters.'
    width, height = page.get_size()
    scale = min(2, 1800 / max(width, height))
    bitmap = page.render(scale=scale)
    image = bitmap.to_pil()
    name = f'page-{page_number}.png'; image.save(Path('/outputs')/name)
    summary['preview_image'] = name
    bitmap.close(); text_page.close(); page.close(); document.close()

assistant.publish(summary)
